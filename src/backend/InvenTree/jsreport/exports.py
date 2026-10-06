import math
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from rest_framework import serializers as drf_serializers
from rest_framework.response import Response

from common.models import DataOutput
from common.serializers import DataOutputSerializer
from InvenTree.helpers import str2bool

MIN_WIDTH = 8
MAX_WIDTH = 40

IMAGE_SIZE = (100, 100)  # max picture size in pixels, aspect ratio is kept
IMAGE_LABEL = "image"  # columns with this label are exported as pictures
IMAGE_PAD = 8  # extra pixels added to the row height around the picture
DEFAULT_ROW_PX = 20  # default Excel row height (15pt) in pixels

THIN_SIDE = Side(style="thin")
DATA_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)
MIDDLE = Alignment(vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Same pale yellow as the on-screen table; header is the same colour at 85% opacity.
DATA_FILL = PatternFill("solid", start_color="FEF9C2", end_color="FEF9C2")
HEAD_FILL = PatternFill("solid", start_color="FEFACB", end_color="FEFACB")


def _resolve(root, relative):
    """Return the file under root, or None (never leaves root)."""
    root = Path(root).resolve()
    candidate = (root / str(relative).lstrip("/")).resolve()

    if root in candidate.parents and candidate.is_file():
        return candidate

    return None


def image_file(value):
    """Return the local file for an image url / path, or None.

    Besides MEDIA_URL / MEDIA_ROOT, extra url prefixes can be mapped to
    folders with a setting, e.g.:

        JSREPORT_IMAGE_ROOTS = {"/costcardimage/": "/data/costcard_images"}
    """
    if not value:
        return None

    text = str(value)
    path = urlparse(text).path if text.startswith(("http://", "https://")) else text

    for prefix, folder in (getattr(settings, "JSREPORT_IMAGE_ROOTS", None) or {}).items():
        if path.startswith(prefix):
            found = _resolve(folder, path[len(prefix):])
            if found:
                return found

    media_url = settings.MEDIA_URL or "/media/"
    if path.startswith(media_url):
        path = path[len(media_url):]

    return _resolve(settings.MEDIA_ROOT, path)


class RowBlock:
    """One record that spans several stacked lines in a grouped sheet.

    Columns that are merged (see GroupedLayout) take their value from the
    first line only.
    """

    def __init__(self, lines):
        self.lines = [list(line) for line in lines] or [[]]


class GroupedLayout:
    """Two-level header layout.

    The headers passed with it are (group, label) tuples:
      * group None  -> the label spans both header rows and the cells of that
        column are merged over all lines of a RowBlock.
      * group "X"   -> the label sits under a merged group heading "X".

    merged_groups  extra groups whose cells are merged over a RowBlock
    filled_groups  groups drawn with the yellow fill
    freeze_cols    number of left columns to keep visible when scrolling
    """

    def __init__(self, merged_groups=(), filled_groups=(), freeze_cols=0):
        self.merged_groups = set(merged_groups)
        self.filled_groups = set(filled_groups)
        self.freeze_cols = freeze_cols


class SimpleSheetBuilder:
    def __init__(self, title="Report"):
        self.title = title
        self._buffers = []  # keeps image buffers alive until the workbook is saved

    def build(self, headers, rows, image_cols=None, layout=None):
        return self.build_multi(
            headers, [(self.title, rows)], image_cols=image_cols, layout=layout
        )

    def build_multi(self, headers, sheets, image_cols=None, layout=None):
        workbook = Workbook()
        workbook.remove(workbook.active)

        for title, rows in sheets or [(self.title, [])]:
            sheet = workbook.create_sheet(self._safe_title(title))

            if layout is not None:
                self._fill_grouped(sheet, headers, rows, image_cols or set(), layout)
            else:
                self._fill(sheet, headers, rows, image_cols or set())

        return workbook

    def to_bytes(self, workbook):
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _safe_title(title):
        return re.sub(r"[\[\]:*?/\\]", "-", str(title)).strip()[:31] or "Sheet"

    # ------------------------------------------------------------------
    # flat sheets
    # ------------------------------------------------------------------
    def _fill(self, ws, headers, rows, image_cols):
        row = 1
        ncols = len(headers)

        for col, head in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=head)
            cell.font = Font(bold=True)
            cell.border = DATA_BORDER

        cell_width = self._image_col_width() * 7 + 5  # characters -> pixels

        for values in rows:
            row += 1
            pictures = []

            for col in range(1, ncols + 1):
                value = values[col - 1] if col <= len(values) else None
                cell = ws.cell(row=row, column=col)
                cell.border = DATA_BORDER

                if col in image_cols:
                    picture = self._load_image(value)
                    if picture:
                        pictures.append((col, picture))
                else:
                    cell.value = value

            if pictures:
                height = max(picture[1].size[1] for _, picture in pictures)
                row_height = height + IMAGE_PAD
                ws.row_dimensions[row].height = row_height * 0.75  # pixels -> points

                for col in range(1, ncols + 1):
                    ws.cell(row=row, column=col).alignment = MIDDLE

                for col, (image, picture) in pictures:
                    self._place_image(ws, image, picture, col, row, cell_width, row_height)

        self._fit_columns(ws, ncols, image_cols)

    # ------------------------------------------------------------------
    # grouped sheets (two-level header, merged blocks, colours)
    # ------------------------------------------------------------------
    def _fill_grouped(self, ws, headers, rows, image_cols, layout):
        head1, head2 = 1, 2
        ncols = len(headers)

        filled = {
            col
            for col, (group, _) in enumerate(headers, start=1)
            if group in layout.filled_groups
        }
        merged = {
            col
            for col, (group, _) in enumerate(headers, start=1)
            if group is None or group in layout.merged_groups
        }

        # ---- two-row header -------------------------------------------
        header_merges = []
        col = 1
        while col <= ncols:
            group, label = headers[col - 1]

            if group is None:
                ws.cell(row=head1, column=col, value=label)
                header_merges.append((head1, col, head2, col))
                span = 1
            else:
                span = 1
                while col + span <= ncols and headers[col + span - 1][0] == group:
                    span += 1

                ws.cell(row=head1, column=col, value=group)
                if span > 1:
                    header_merges.append((head1, col, head1, col + span - 1))

                for sub in range(col, col + span):
                    ws.cell(row=head2, column=sub, value=headers[sub - 1][1])

            col += span

        for r in (head1, head2):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(bold=True)
                cell.border = DATA_BORDER
                cell.alignment = CENTER
                if c in filled:
                    cell.fill = HEAD_FILL

        for r1, c1, r2, c2 in header_merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

        # ---- data blocks ----------------------------------------------
        cell_width = self._image_col_width() * 7 + 5  # characters -> pixels
        row = head2

        for block in rows:
            lines = block.lines if isinstance(block, RowBlock) else [block]
            count = len(lines)
            first = row + 1
            pictures = []

            for index, values in enumerate(lines):
                row += 1

                for col in range(1, ncols + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.border = DATA_BORDER
                    cell.alignment = MIDDLE
                    if col in filled:
                        cell.fill = DATA_FILL

                    value = values[col - 1] if col <= len(values) else None

                    if col in image_cols:
                        if index == 0:
                            picture = self._load_image(value)
                            if picture:
                                pictures.append((col, picture))
                    elif col in merged and index > 0:
                        continue  # covered by the merge, value comes from line 1
                    else:
                        cell.value = value

            last = row

            # make the block tall enough for the tallest picture
            line_px = DEFAULT_ROW_PX
            if pictures:
                tallest = max(picture[1].size[1] for _, picture in pictures)
                line_px = max(DEFAULT_ROW_PX, math.ceil((tallest + IMAGE_PAD) / count))
                for r in range(first, last + 1):
                    ws.row_dimensions[r].height = line_px * 0.75  # pixels -> points

            block_px = line_px * count

            for col, (image, picture) in pictures:
                width, height = picture.size
                top = max(0, (block_px - height) / 2)
                offset_row = min(int(top // line_px), count - 1)
                self._place_image_at(
                    ws,
                    image,
                    picture,
                    col,
                    first + offset_row,
                    col_off=(cell_width - width) / 2,
                    row_off=top - offset_row * line_px,
                )

            if count > 1:
                for col in sorted(merged):
                    ws.merge_cells(
                        start_row=first, start_column=col, end_row=last, end_column=col
                    )

        self._fit_grouped_columns(ws, headers, head2, image_cols)

        if layout.freeze_cols:
            ws.freeze_panes = ws.cell(row=head2 + 1, column=layout.freeze_cols + 1)

    # ------------------------------------------------------------------
    # images
    # ------------------------------------------------------------------
    @staticmethod
    def _image_col_width():
        return IMAGE_SIZE[0] / 7 + 2  # pixels -> characters

    def _load_image(self, value):
        """Return (excel image, resized picture) for a url / path, or None."""
        path = image_file(value)

        if path is None:
            return None

        try:
            picture = PILImage.open(path)
            picture.thumbnail(IMAGE_SIZE)

            if picture.mode not in ("RGB", "RGBA"):
                picture = picture.convert("RGB")

            buffer = BytesIO()
            picture.save(buffer, format="PNG")
            buffer.seek(0)
            self._buffers.append(buffer)

            return XLImage(buffer), picture
        except Exception:
            return None

    @staticmethod
    def _place_image(ws, image, picture, col, row, cell_width, row_height):
        """Anchor the picture in the middle of the cell, both ways."""
        width, height = picture.size
        SimpleSheetBuilder._place_image_at(
            ws,
            image,
            picture,
            col,
            row,
            col_off=(cell_width - width) / 2,
            row_off=(row_height - height) / 2,
        )

    @staticmethod
    def _place_image_at(ws, image, picture, col, row, col_off=0, row_off=0):
        """Anchor the picture at a pixel offset inside the given cell."""
        width, height = picture.size
        marker = AnchorMarker(
            col=col - 1,
            row=row - 1,
            colOff=pixels_to_EMU(max(0, col_off)),
            rowOff=pixels_to_EMU(max(0, row_off)),
        )
        image.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
        )
        ws.add_image(image)

    # ------------------------------------------------------------------
    # column widths
    # ------------------------------------------------------------------
    @staticmethod
    def _fit_columns(ws, ncols, image_cols=()):
        widths = {}

        for row in ws.iter_rows(max_col=ncols):
            for cell in row:
                if cell.value is not None:
                    length = len(str(cell.value))
                    widths[cell.column] = max(widths.get(cell.column, 0), length)

        for col in range(1, ncols + 1):
            if col in image_cols:
                width = SimpleSheetBuilder._image_col_width()
            else:
                width = min(max(widths.get(col, 0) + 2, MIN_WIDTH), MAX_WIDTH)
            ws.column_dimensions[get_column_letter(col)].width = width

    @staticmethod
    def _fit_grouped_columns(ws, headers, head2, image_cols):
        """Widths from the sub-headers and data only.

        The merged group headings are left out so they do not stretch the
        first column underneath them.
        """
        ncols = len(headers)
        widths = {}

        for row in ws.iter_rows(min_row=head2, max_col=ncols):
            for cell in row:
                if cell.value is not None:
                    length = len(str(cell.value))
                    widths[cell.column] = max(widths.get(cell.column, 0), length)

        for col, (group, label) in enumerate(headers, start=1):
            if group is None:
                widths[col] = max(widths.get(col, 0), len(str(label)))

        for col in range(1, ncols + 1):
            if col in image_cols:
                width = SimpleSheetBuilder._image_col_width()
            else:
                width = min(max(widths.get(col, 0) + 2, MIN_WIDTH), MAX_WIDTH)
            ws.column_dimensions[get_column_letter(col)].width = width


class JSReportExportMixin:
    export_name = "Report"
    export_plugin = "jsreport"
    export_serial = False
    export_image_fields = None  # field names to export as pictures; None = label "Image"

    def get_export_sheet_title(self, instance):
        return None

    def get_export_table(self, queryset, items, instances):
        """Optional custom table.

        Return None for the default one-row-per-record export, or a tuple
        (headers, rows, layout) where headers are (group, label) tuples, rows
        holds one RowBlock per instance (same order as `instances`) and layout
        is a GroupedLayout. With export_serial = True, column 1 must be the
        "#" column; it is numbered here (restarting on every sheet).
        """
        return None

    def list(self, request, *args, **kwargs):
        params = request.query_params

        if str2bool(params.get("export")) or params.get("export_format"):
            return self.export_response(request)

        return super().list(request, *args, **kwargs)

    def export_response(self, request):
        export_format = request.query_params.get("export_format") or "xlsx"

        if export_format.lower() != "xlsx":
            return Response({"detail": "Only xlsx export is supported."}, status=400)

        queryset = self.filter_queryset(self.get_queryset())
        serializer = self.get_serializer(queryset, many=True)
        fields = serializer.child.fields
        items = serializer.data
        instances = list(queryset)

        layout = None
        table = self.get_export_table(queryset, items, instances)

        if table is not None:
            headers, rows, layout = table
            labels = [label for _, label in headers]
            image_cols = {
                col
                for col, label in enumerate(labels, start=1)
                if str(label).strip().lower() == IMAGE_LABEL
            }
            sheets = self._split_sheets(instances, rows)

            if self.export_serial:
                for _, sheet_rows in sheets:
                    for number, block in enumerate(sheet_rows, start=1):
                        block.lines[0][0] = number
        else:
            names = list(fields.keys())
            headers = [fields[name].label or name for name in names]
            rows = [
                [self._cell(fields[name], item.get(name)) for name in names]
                for item in items
            ]

            image_cols = self._image_cols(names, headers)
            sheets = self._split_sheets(instances, rows)

            if self.export_serial:
                headers.insert(0, "#")
                image_cols = {col + 1 for col in image_cols}
                sheets = [
                    (title, [[number] + row for number, row in enumerate(sheet_rows, start=1)])
                    for title, sheet_rows in sheets
                ]

        builder = SimpleSheetBuilder(self.export_name.replace("_", " "))
        content = builder.to_bytes(
            builder.build_multi(
                headers,
                sheets,
                image_cols=image_cols,
                layout=layout,
            )
        )

        output = DataOutput.objects.create(
            user=request.user if request.user.is_authenticated else None,
            total=len(rows),
            progress=100,
            complete=True,
            output_type=DataOutput.DataOutputTypes.EXPORT,
            plugin=self.export_plugin,
        )

        timestamp = timezone.now().strftime("%Y_%m_%d_%H_%M")
        output.output.save(
            f"{self.export_name}_{timestamp}.xlsx",
            ContentFile(content),
            save=True,
        )

        return Response(DataOutputSerializer(output).data, status=200)

    def _image_cols(self, names, headers):
        """1-based column numbers that hold pictures."""
        if self.export_image_fields is not None:
            return {names.index(n) + 1 for n in self.export_image_fields if n in names}

        return {
            col
            for col, head in enumerate(headers, start=1)
            if str(head).strip().lower() == IMAGE_LABEL
        }

    def _split_sheets(self, instances, rows):
        titles = [self.get_export_sheet_title(instance) for instance in instances]

        if not any(titles):
            return [(self.export_name.replace("_", " "), rows)]

        grouped = {}
        for title, row in zip(titles, rows):
            grouped.setdefault(title or "No Vendor", []).append(row)

        ordered = sorted(grouped, key=lambda t: t.casefold().replace(" ", ""))

        return [
            (f"{number}-{title}", grouped[title])
            for number, title in enumerate(ordered, start=1)
        ]

    @staticmethod
    def _cell(field, value):
        if value is None:
            return None
        if isinstance(field, drf_serializers.DecimalField):
            return float(value)
        return value