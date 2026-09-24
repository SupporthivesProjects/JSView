import json
from decimal import Decimal, InvalidOperation
from io import BytesIO
from urllib.request import Request, urlopen

from django.db import transaction

from rest_framework import serializers as drf_serializers

from InvenTree.serializers import InvenTreeModelSerializer

from data_exporter.mixins import DataExportSerializerMixin
from importer.mixins import DataImportSerializerMixin
from importer.registry import register_importer

from revision.costcard.services import create_cost_card_version

from .models import (
    CostCard,
    CostCardColorStoneLine,
    CostCardDiamondLine,
    CostCardFinishLine,
    StonePlace,
)


@register_importer()
class StonePlaceSerializer(
    DataImportSerializerMixin,
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    class Meta:
        model = StonePlace
        fields = [
            'pk',
            'name',
            'description',
            'active',
            'created_at',
            'updated_at',
        ]


@register_importer()
class CostCardDiamondLineSerializer(
    DataImportSerializerMixin,
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    class Meta:
        model = CostCardDiamondLine
        fields = [
            'pk',
            'cost_card',
            'stone',
            'shape',
            'mm_size',
            'color',
            'cut',
            'quality',
            'setting',
            'stone_place',
            'rate_source',
            'pointer',
            'sieve_size',
            'pcs',
            'cts',
            'default_rate',
            'pc',
            'rate',
            'amount',
            'labour_rate',
            'labour_amount',
            'active',
            'created_at',
            'updated_at',
        ]


@register_importer()
class CostCardColorStoneLineSerializer(
    DataImportSerializerMixin,
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    class Meta:
        model = CostCardColorStoneLine
        fields = [
            'pk',
            'cost_card',
            'stone',
            'shape',
            'mm_size',
            'color',
            'cut',
            'quality',
            'setting',
            'stone_place',
            'rate_source',
            'pointer',
            'sieve_size',
            'pcs',
            'cts',
            'default_rate',
            'pc',
            'rate',
            'amount',
            'labour_rate',
            'labour_amount',
            'active',
            'created_at',
            'updated_at',
        ]


@register_importer()
class CostCardFinishLineSerializer(
    DataImportSerializerMixin,
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    class Meta:
        model = CostCardFinishLine
        fields = [
            'pk',
            'cost_card',
            'finish_type',
            'rate',
            'active',
            'created_at',
            'updated_at',
        ]


class NestedDiamondLineSerializer(drf_serializers.ModelSerializer):
    id = drf_serializers.IntegerField(
        required=False,
        allow_null=True,
    )

    class Meta:
        model = CostCardDiamondLine
        fields = [
            'id',
            'stone',
            'shape',
            'mm_size',
            'color',
            'cut',
            'quality',
            'setting',
            'stone_place',
            'rate_source',
            'pointer',
            'sieve_size',
            'pcs',
            'cts',
            'default_rate',
            'pc',
            'rate',
            'amount',
            'labour_rate',
            'labour_amount',
            'active',
        ]


class NestedColorStoneLineSerializer(drf_serializers.ModelSerializer):
    id = drf_serializers.IntegerField(
        required=False,
        allow_null=True,
    )

    class Meta:
        model = CostCardColorStoneLine
        fields = [
            'id',
            'stone',
            'shape',
            'mm_size',
            'color',
            'cut',
            'quality',
            'setting',
            'stone_place',
            'rate_source',
            'pointer',
            'sieve_size',
            'pcs',
            'cts',
            'default_rate',
            'pc',
            'rate',
            'amount',
            'labour_rate',
            'labour_amount',
            'active',
        ]


class NestedFinishLineSerializer(drf_serializers.ModelSerializer):
    id = drf_serializers.IntegerField(
        required=False,
        allow_null=True,
    )

    class Meta:
        model = CostCardFinishLine
        fields = [
            'id',
            'finish_type',
            'rate',
            'active',
        ]


@register_importer()
class CostCardSerializer(
    DataImportSerializerMixin,
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    diamond_lines = NestedDiamondLineSerializer(
        many=True,
        required=False,
    )

    colorstone_lines = NestedColorStoneLineSerializer(
        many=True,
        required=False,
    )

    finish_lines = NestedFinishLineSerializer(
        many=True,
        required=False,
    )

    NESTED_LINE_FIELDS = (
        'diamond_lines',
        'colorstone_lines',
        'finish_lines',
    )

    def skip_create_fields(self):
        return [*super().skip_create_fields(), *self.NESTED_LINE_FIELDS]

    class Meta:
        model = CostCard
        fields = [
            'pk',
            'cost_card_no',
            'our_style_no',
            'vendor_style_no',
            'vendor',
            'customer',
            'category',
            'sub_category',
            'metal_purity',
            'karat',
            'metal_grams',
            'finding_type',
            'finding_item',
            'finding_price',
            'gross_weight',
            'net_weight',
            'troy_ounce_price',
            'height_mm',
            'height_inch',
            'length_mm',
            'length_inch',
            'width_mm',
            'width_inch',
            'shank_size_mm',
            'shank_size_inch',
            'drape_length_mm',
            'drape_length_inch',
            'design_note',
            'special_note',
            'remarks',
            'front_view',
            'side_view',
            'back_view',
            'labour_finish_amount',
            'labour_diamond_amount',
            'labour_colorstone_amount',
            'metal_loss_pct',
            'metal_loss_amount',
            'metal_amount',
            'dia_pcs',
            'dia_cts',
            'dia_amount',
            'col_pcs',
            'col_cts',
            'col_amount',
            'stone_pcs',
            'stone_cts',
            'stone_amount',
            'labour_amount',
            'dia_handling_pct',
            'dia_handling_amount',
            'col_handling_pct',
            'col_handling_amount',
            'vendor_markup_pct',
            'vendor_markup_amount',
            'fob',
            'duty_pct',
            'duty_amount',
            'margin_pct',
            'margin_amount',
            'final_amount',
            'remarks_full',
            'diamond_lines',
            'colorstone_lines',
            'finish_lines',
            'active',
            'created_at',
            'updated_at',
        ]

        read_only_fields = [
            'pk',
            'cost_card_no',
            'front_view',
            'side_view',
            'back_view',
            'created_at',
            'updated_at',
        ]

    def _current_user(self):
        request = self.context.get('request')
        user = getattr(request, 'user', None) if request else None

        if user is not None and getattr(user, 'is_authenticated', False):
            return user

        return None

    @staticmethod
    def _normalize_nested_payload(value):
        if value is None:
            return None

        if isinstance(value, (list, tuple)) and len(value) == 1 and isinstance(value[0], str):
            value = value[0]

        if isinstance(value, str):
            try:
                value = json.loads(value)
            except (TypeError, ValueError):
                return value

        if isinstance(value, dict):
            value = [value]

        return value

    def to_internal_value(self, data):
        if hasattr(data, 'copy'):
            data = data.copy()
        else:
            data = dict(data)

        diamond_lines = self._normalize_nested_payload(
            data.pop('diamond_lines', None)
        )
        colorstone_lines = self._normalize_nested_payload(
            data.pop('colorstone_lines', None)
        )
        finish_lines = self._normalize_nested_payload(
            data.pop('finish_lines', None)
        )

        validated_data = super().to_internal_value(data)

        if diamond_lines is not None:
            validated_data['diamond_lines'] = NestedDiamondLineSerializer(
                many=True
            ).to_internal_value(diamond_lines)

        if colorstone_lines is not None:
            validated_data['colorstone_lines'] = NestedColorStoneLineSerializer(
                many=True
            ).to_internal_value(colorstone_lines)

        if finish_lines is not None:
            validated_data['finish_lines'] = NestedFinishLineSerializer(
                many=True
            ).to_internal_value(finish_lines)

        if self.instance is None:
            validated_data['cost_card_no'] = self._generate_cost_card_no()

        return validated_data

    def _generate_cost_card_no(self):
        last_card = CostCard.objects.order_by('-id').first()

        if last_card and last_card.cost_card_no:
            try:
                number = int(
                    ''.join(
                        char
                        for char in last_card.cost_card_no
                        if char.isdigit()
                    )
                )
                return f"{number + 1:05d}"
            except (ValueError, TypeError):
                pass

        return "00001"

    def _scalar_fields(self, validated_data):
        return {
            key: value
            for key, value in validated_data.items()
            if key not in self.NESTED_LINE_FIELDS
        }

    @transaction.atomic
    def create(self, validated_data):
        diamond_lines = validated_data.pop('diamond_lines', [])
        colorstone_lines = validated_data.pop('colorstone_lines', [])
        finish_lines = validated_data.pop('finish_lines', [])

        validated_data['cost_card_no'] = self._generate_cost_card_no()

        cost_card = CostCard.objects.create(
            **self._scalar_fields(validated_data)
        )

        for line_data in diamond_lines:
            line_data.pop('id', None)

            CostCardDiamondLine.objects.create(
                cost_card=cost_card,
                **line_data,
            )

        for line_data in colorstone_lines:
            line_data.pop('id', None)

            CostCardColorStoneLine.objects.create(
                cost_card=cost_card,
                **line_data,
            )

        for line_data in finish_lines:
            line_data.pop('id', None)

            CostCardFinishLine.objects.create(
                cost_card=cost_card,
                **line_data,
            )

        create_cost_card_version(
            cost_card=cost_card,
            user=self._current_user(),
        )

        return cost_card

    @transaction.atomic
    def update(self, instance, validated_data):
        diamond_lines = validated_data.pop('diamond_lines', None)
        colorstone_lines = validated_data.pop('colorstone_lines', None)
        finish_lines = validated_data.pop('finish_lines', None)

        validated_data.pop('cost_card_no', None)

        for attr, value in self._scalar_fields(validated_data).items():
            setattr(instance, attr, value)

        instance.save()

        if diamond_lines is not None:
            self._sync_lines(
                instance,
                CostCardDiamondLine,
                diamond_lines,
                True,
            )

        if colorstone_lines is not None:
            self._sync_lines(
                instance,
                CostCardColorStoneLine,
                colorstone_lines,
                True,
            )

        if finish_lines is not None:
            self._sync_lines(
                instance,
                CostCardFinishLine,
                finish_lines,
                True,
            )

        create_cost_card_version(
            cost_card=instance,
            user=self._current_user(),
        )

        return instance

    @staticmethod
    def _sync_lines(
        cost_card,
        model_cls,
        lines_data,
        allow_delete,
    ):
        existing = {
            obj.pk: obj
            for obj in model_cls.objects.filter(
                cost_card=cost_card
            )
        }

        seen_ids = set()

        for data in lines_data:
            line_id = data.pop('id', None)
            data.pop('cost_card', None)

            if line_id is not None:
                line_instance = existing.get(line_id)

                if line_instance is None:
                    raise drf_serializers.ValidationError(
                        {
                            'id': (
                                f'{model_cls.__name__} with id '
                                f'{line_id} does not belong to '
                                f'CostCard {cost_card.pk}.'
                            )
                        }
                    )

                for attr, value in data.items():
                    setattr(line_instance, attr, value)

                line_instance.save()

                seen_ids.add(line_id)

            else:
                model_cls.objects.create(
                    cost_card=cost_card,
                    **data,
                )

        if allow_delete:
            for line_id, line_instance in existing.items():
                if line_id not in seen_ids:
                    line_instance.delete()


class CostCardImageSerializer(InvenTreeModelSerializer):
    class Meta:
        model = CostCard
        fields = [
            'pk',
            'front_view',
            'side_view',
            'back_view',
        ]

    def _current_user(self):
        request = self.context.get('request')
        user = getattr(request, 'user', None) if request else None

        if user is not None and getattr(user, 'is_authenticated', False):
            return user

        return None

    def update(self, instance, validated_data):
        instance = super().update(instance, validated_data)

        create_cost_card_version(
            cost_card=instance,
            user=self._current_user(),
        )

        return instance


# def _dec(value):
#     if value in (None, '', 'null'):
#         return None
#     try:
#         return Decimal(str(value))
#     except InvalidOperation:
#         return None


def _dec(value):
    if isinstance(value, (list, tuple)):
        value = value[0] if value else None

    if value in (None, '', 'null'):
        return None

    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        return None


class CostCardPicturePresentationSerializer(
    DataExportSerializerMixin,
    InvenTreeModelSerializer,
):
    picture = drf_serializers.SerializerMethodField(label='Picture')
    our_style_no = drf_serializers.CharField(label='JS Style No.')
    vendor_style_no = drf_serializers.CharField(
        label='Vendor Style No.',
        allow_null=True,
    )
    kt = drf_serializers.SerializerMethodField(label='KT')
    dia_cts = drf_serializers.DecimalField(
        max_digits=10,
        decimal_places=4,
        label='DIA CTS.',
    )
    quality = drf_serializers.SerializerMethodField(label='QUALITY')
    col_cts = drf_serializers.DecimalField(
        max_digits=10,
        decimal_places=4,
        label='COL CTS.',
    )
    cost = drf_serializers.SerializerMethodField(label='COST $')
    comments = drf_serializers.CharField(
        source='remarks',
        label='COMMENTS',
        allow_null=True,
    )
    duty = drf_serializers.SerializerMethodField(label='Duty %')
    margin = drf_serializers.SerializerMethodField(label='Margin %')

    class Meta:
        model = CostCard
        fields = [
            'picture',
            'our_style_no',
            'vendor_style_no',
            'kt',
            'dia_cts',
            'quality',
            'col_cts',
            'cost',
            'comments',
            'duty',
            'margin',
        ]

    # def _param(self, key):
    #     value = self.context.get(key)
    #     if value in (None, ''):
    #         request = self.context.get('request')
    #         params = getattr(request, 'query_params', None) or {}
    #         value = params.get(key)
    #     return value

    def _param(self, key):
        value = self.context.get(key)

        if isinstance(value, (list, tuple)):
            value = value[0] if value else None

        if value in (None, ''):
            request = self.context.get('request')
            params = getattr(request, 'query_params', None) or {}
            value = params.get(key)

            if isinstance(value, (list, tuple)):
                value = value[0] if value else None

        return value

    def _duty_pct(self, obj):
        value = _dec(self._param('duty_pct'))
        return value if value is not None else obj.duty_pct

    def _margin_pct(self, obj):
        value = _dec(self._param('margin_pct'))
        return value if value is not None else obj.margin_pct

    # def get_picture(self, obj):
    #     if not obj.front_view:
    #         return ''
    #     url = obj.front_view.url
    #     request = self.context.get('request')
    #     return request.build_absolute_uri(url) if request else url

    def get_picture(self, obj):
        if not obj.front_view:
            return ''

        url = obj.front_view.url
        request = self.context.get('request')

        build = getattr(request, 'build_absolute_uri', None)

        if build:
            return build(url)

        try:
            from InvenTree.helpers_model import construct_absolute_url

            return construct_absolute_url(url)
        except Exception:
            return url

    def get_kt(self, obj):
        return obj.metal_purity.name if obj.metal_purity else (obj.karat or '')

    def get_quality(self, obj):
        line = obj.diamond_lines.first()
        return line.stone.name if line and line.stone else ''

    def _metal_ratio(self, obj):
        name = (
            obj.metal_purity.name
            if obj.metal_purity
            else ''
        ).lower()

        key = (
            'silver_troy_ounce'
            if 'silver' in name
            else 'gold_troy_ounce'
        )

        new_price = _dec(self._param(key))

        if new_price is None or not obj.troy_ounce_price:
            return None

        return new_price / obj.troy_ounce_price

    def _fob(self, obj):
        ratio = self._metal_ratio(obj)

        if ratio is None:
            return obj.fob

        return obj.fob + (
            obj.metal_amount + obj.metal_loss_amount
        ) * (ratio - 1)

    def get_cost(self, obj):
        keys = (
            'duty_pct',
            'margin_pct',
            'gold_troy_ounce',
            'silver_troy_ounce',
        )

        if all(_dec(self._param(k)) is None for k in keys):
            return obj.final_amount

        fob = self._fob(obj)

        duty_amt = (
            fob
            * self._duty_pct(obj)
            / 100
        )

        margin_amt = (
            (fob + duty_amt)
            * self._margin_pct(obj)
            / 100
        )

        return (
            fob
            + duty_amt
            + margin_amt
        ).quantize(Decimal('0.01'))

    def get_duty(self, obj):
        return self._duty_pct(obj)

    def get_margin(self, obj):
        return self._margin_pct(obj)

    def export_to_file(self, data, headers, file_format):
        """
        Keep the existing InvenTree export behavior for all formats
        except XLSX.

        For XLSX, generate the workbook here so the Picture column
        contains the actual embedded image instead of the image URL.
        """

        if file_format != 'xlsx':
            return super().export_to_file(
                data,
                headers,
                file_format,
            )

        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Picture Presentation'

        field_names = list(headers.keys())
        field_headers = list(headers.values())

        # Header row
        for column_index, header in enumerate(field_headers, start=1):
            worksheet.cell(
                row=1,
                column=column_index,
                value=header,
            )

        # Find the Picture column
        picture_column = None

        for index, field_name in enumerate(field_names, start=1):
            if field_name == 'picture':
                picture_column = index
                break

        # Existing data
        for row_index, row in enumerate(data, start=2):
            for column_index, field_name in enumerate(
                field_names,
                start=1,
            ):
                value = self.get_nested_value(
                    row,
                    field_name,
                )

                # Do not write the image URL into Excel.
                # The actual image will be embedded below.
                if (
                    picture_column is not None
                    and column_index == picture_column
                ):
                    continue

                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

            # Embed actual image
            if picture_column is not None:
                picture_url = self.get_nested_value(
                    row,
                    'picture',
                )

                if picture_url:
                    try:
                        request = Request(
                            picture_url,
                            headers={
                                'User-Agent': 'Mozilla/5.0',
                            },
                        )

                        with urlopen(
                            request,
                            timeout=15,
                        ) as response:
                            image_data = response.read()

                        image_stream = BytesIO(image_data)

                        excel_image = ExcelImage(
                            image_stream
                        )

                        # Keep image size reasonable
                        max_width = 120
                        max_height = 120

                        if excel_image.width > max_width:
                            ratio = (
                                max_width
                                / excel_image.width
                            )
                            excel_image.width = max_width
                            excel_image.height = int(
                                excel_image.height * ratio
                            )

                        if excel_image.height > max_height:
                            ratio = (
                                max_height
                                / excel_image.height
                            )
                            excel_image.height = max_height
                            excel_image.width = int(
                                excel_image.width * ratio
                            )

                        cell_reference = (
                            f'{get_column_letter(picture_column)}'
                            f'{row_index}'
                        )

                        worksheet.add_image(
                            excel_image,
                            cell_reference,
                        )

                        # Make row high enough for image
                        worksheet.row_dimensions[
                            row_index
                        ].height = 95

                    except Exception:
                        # If image cannot be downloaded,
                        # keep the rest of the export working.
                        worksheet.cell(
                            row=row_index,
                            column=picture_column,
                            value='',
                        )

        # Basic column sizing
        for column_index, field_name in enumerate(
            field_names,
            start=1,
        ):
            column_letter = get_column_letter(
                column_index
            )

            if (
                picture_column is not None
                and column_index == picture_column
            ):
                worksheet.column_dimensions[
                    column_letter
                ].width = 22
                continue

            max_length = len(
                str(
                    worksheet.cell(
                        row=1,
                        column=column_index,
                    ).value or ''
                )
            )

            for row_index in range(
                2,
                worksheet.max_row + 1,
            ):
                value = worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value

                if value is not None:
                    max_length = max(
                        max_length,
                        len(str(value)),
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                35,
            )

        output = BytesIO()

        workbook.save(output)

        return output.getvalue()