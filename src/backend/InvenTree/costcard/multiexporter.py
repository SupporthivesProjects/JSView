import os
import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

DEFAULT_LOGO_PATH = os.path.join(settings.STATIC_ROOT, 'logos', 'logo_for_excel_sheet.png')

FOOTER_TEXT = 'The above quotation is based on current market price and is subject to change at any time without prior notice'
LAST_COL = 16
THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
MONEY = '0.00'
MIN_COL_WIDTH = 4
MAX_COL_WIDTH = 26
MIN_FINISH_ROWS = 2
ZERO = Decimal('0')
CENT = Decimal('0.01')
HUNDRED = Decimal('100')
STUDDING_HEADERS = [
    'Type', 'Shape', 'Cut', 'MM Size', 'Sieve Size', 'Stone Type', 'Colour', 'Pointer',
    'Pcs', 'Carats', 'P/C', 'Rate', 'Amount', 'Setting', 'Rate', 'Amount',
]
STUDDING_FORMATS = {10: MONEY, 12: MONEY, 13: MONEY, 15: MONEY, 16: MONEY}
COST_HEADERS = ['Metal', 'Studding', 'Labour', 'Markup %', 'With Markup %', 'F O B', 'Duty %', 'With Duty', 'Margin %', 'Final Price']
COST_KEYS = ['metal', 'studding', 'labour', 'markup_pct', 'markup_amount', 'fob', 'duty_pct', 'with_duty', 'margin_pct', 'final']
COST_COLUMNS = ['F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O']


def dec(value):
    if isinstance(value, (list, tuple)):
        value = value[0] if value else None
    if value in (None, '', 'null'):
        return None
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        return None


def num(value):
    return None if value is None else float(value)


def money(value):
    return Decimal(value).quantize(CENT, rounding=ROUND_HALF_UP)


def percent_of(base, pct):
    return money(Decimal(base) * Decimal(pct or 0) / HUNDRED)


def total(values):
    return sum(values, ZERO)


def name_of(value):
    return '' if value is None else str(getattr(value, 'name', value))


def size_of(value):
    return '' if value is None else str(value.mm_size)


def put(ws, ref, value, bold=False, border=True, fmt=None, size=11):
    cell = ws[ref]
    cell.value = value
    cell.font = Font(bold=bold, size=size)
    cell.alignment = CENTER
    if border:
        cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def merge(ws, cell_range, value, bold=False, border=True, fmt=None, size=11):
    ws.merge_cells(cell_range)
    put(ws, cell_range.split(':')[0], value, bold, border, fmt, size)
    if border:
        for row in ws[cell_range]:
            for cell in row:
                cell.border = BORDER


class CostCardSheetBuilder:
    SECTIONS = ('title', 'header', 'metal', 'studding', 'labour_and_cost', 'design', 'footer')

    def __init__(self, overrides=None, modified_by='', logo_path=None, footer_text=FOOTER_TEXT, sheet_title=None, sections=None):
        self.overrides = {key: dec(value) for key, value in (overrides or {}).items()}
        self.modified_by = modified_by
        self.logo_path = logo_path or DEFAULT_LOGO_PATH
        self.footer_text = footer_text
        self.sheet_title = sheet_title
        self.sections = sections or self.SECTIONS
        self._buffers = []

    def build(self, cards):
        workbook = Workbook()
        workbook.remove(workbook.active)
        used = set()
        for card in cards:
            self.fill_sheet(workbook.create_sheet(self._unique_title(card, used)), card)
        if not workbook.sheetnames:
            workbook.create_sheet('No data')
        return workbook

    def to_bytes(self, workbook):
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def fill_sheet(self, ws, card):
        ws.sheet_view.showGridLines = True
        figures = self.figures(card)
        row = 1
        for name in self.sections:
            row = getattr(self, f'_section_{name}')(ws, card, figures, row)
        self._apply_outer_border(ws, row - 1)
        self._autofit_columns(ws)

    def _pct(self, key, stored):
        value = self.overrides.get(key)
        return Decimal(stored or 0) if value is None else value

    def figures(self, card):
        overrides = self.overrides
        purity = card.metal_purity
        text = f"{name_of(getattr(purity, 'metal_type', None))} {name_of(purity)}".lower()
        price = overrides.get('silver_troy_ounce' if 'silver' in text else 'gold_troy_ounce')
        ratio = price / card.troy_ounce_price if price is not None and card.troy_ounce_price else None

        metal_amount, tr_oz = Decimal(card.metal_amount or 0), card.troy_ounce_price
        if ratio is not None:
            metal_amount, tr_oz = metal_amount * ratio, price
        metal_amount = money(metal_amount)

        diamond_lines = list(card.diamond_lines.all())
        color_lines = list(card.colorstone_lines.all())
        finish_lines = list(card.finish_lines.all())
        stone_lines = diamond_lines + color_lines

        if stone_lines:
            diamond_amount = money(total(line.amount for line in diamond_lines))
            color_amount = money(total(line.amount for line in color_lines))
            stone_amount = diamond_amount + color_amount
        else:
            diamond_amount = Decimal(card.dia_amount or 0)
            color_amount = Decimal(card.col_amount or 0)
            stone_amount = Decimal(card.stone_amount or 0)

        if stone_lines or finish_lines:
            labour = money(
                total(line.rate for line in finish_lines)
                + total(line.labour_amount for line in stone_lines)
            )
        else:
            labour = Decimal(card.labour_amount or 0)

        handling = (
            percent_of(diamond_amount, card.dia_handling_pct)
            + percent_of(color_amount, card.col_handling_pct)
        )
        metal = metal_amount + money(card.finding_price or 0)
        studding = stone_amount + handling

        markup_pct = self._pct('markup_pct', card.vendor_markup_pct)
        duty_pct = self._pct('duty_pct', card.duty_pct)
        margin_pct = self._pct('margin_pct', card.margin_pct)

        subtotal = metal + studding + labour
        markup_amount = percent_of(subtotal, markup_pct)
        fob = subtotal + markup_amount
        with_duty = fob + percent_of(fob, duty_pct)
        final = with_duty + percent_of(with_duty, margin_pct)

        return {
            'tr_oz': tr_oz,
            'metal_amount': metal_amount,
            'metal': metal,
            'studding': studding,
            'labour': labour,
            'markup_pct': markup_pct,
            'markup_amount': markup_amount,
            'fob': fob,
            'duty_pct': duty_pct,
            'with_duty': with_duty,
            'margin_pct': margin_pct,
            'final': final,
        }

    def _section_title(self, ws, card, figures, row):
        merge(ws, f'A{row}:P{row}', 'COST BREAKDOWN SHEET', bold=True, border=False, size=14)
        if self.logo_path:
            self._image(ws, self.logo_path, f'B{row + 1}', (230, 90))
            return row + 5
        return row + 1

    def _section_header(self, ws, card, figures, row):
        left = [
            ('CC No :', card.cost_card_no),
            ('Customer :', name_of(card.customer)),
            ('Vendor :', name_of(card.vendor)),
            ('V. Style :', card.vendor_style_no or ''),
        ]
        right = [
            ('Style :', card.our_style_no),
            ('Category :', name_of(card.category)),
            ('Sub Category :', name_of(card.sub_category)),
            ('Date :', timezone.now().strftime('%d %b %Y')),
            ('Modified By :', self.modified_by),
        ]
        for offset, (label, value) in enumerate(left):
            put(ws, f'A{row + offset}', label)
            merge(ws, f'B{row + offset}:C{row + offset}', value)
        for offset, (label, value) in enumerate(right):
            put(ws, f'E{row + offset}', label)
            merge(ws, f'F{row + offset}:G{row + offset}', value)
            ws.row_dimensions[row + offset].height = 22
        if card.front_view:
            self._image(ws, Path(settings.MEDIA_ROOT) / card.front_view.name, f'I{row}', (120, 120))
        return row + 6

    def _section_metal(self, ws, card, figures, row):
        for col, title in zip('ABCDE', ['Type', 'Tr. Oz', 'KT', 'Net Wt.', 'Loss %']):
            put(ws, f'{col}{row}', title, bold=True)
        merge(ws, f'F{row}:G{row}', 'Metal Amount', bold=True)

        purity = card.metal_purity
        values = [
            name_of(getattr(purity, 'metal_type', None)),
            num(figures['tr_oz']),
            card.karat,
            num(card.net_weight),
            num(card.metal_loss_pct),
        ]
        for col, value in zip('ABCDE', values):
            put(ws, f'{col}{row + 1}', value)
        merge(ws, f'F{row + 1}:G{row + 1}', num(figures['metal_amount']), fmt=MONEY)
        return row + 3

    def _section_studding(self, ws, card, figures, row):
        merge(ws, f'A{row}:M{row}', 'Studding Details', bold=True)
        merge(ws, f'N{row}:P{row}', 'Labour', bold=True)
        for col, title in enumerate(STUDDING_HEADERS, start=1):
            put(ws, f'{get_column_letter(col)}{row + 1}', title, bold=True)
        row += 2

        lines = self._lines(card)
        for kind, line in lines:
            values = [
                kind, name_of(line.shape), name_of(line.cut), size_of(line.mm_size), line.sieve_size or '',
                name_of(line.stone), name_of(line.color), num(line.pointer), line.pcs, num(line.cts),
                line.pc, num(line.rate), num(line.amount), name_of(line.setting), num(line.labour_rate), num(line.labour_amount),
            ]
            for col, value in enumerate(values, start=1):
                put(ws, f'{get_column_letter(col)}{row}', value, fmt=STUDDING_FORMATS.get(col))
            row += 1

        has_lines = bool(lines)
        totals = {
            8: 'Total',
            9: sum(line.pcs for _, line in lines) if has_lines else None,
            10: num(total(line.cts for _, line in lines)) if has_lines else None,
            13: num(total(line.amount for _, line in lines)) if has_lines else None,
            16: num(total(line.labour_amount for _, line in lines)) if has_lines else None,
        }
        for col in range(1, LAST_COL + 1):
            put(ws, f'{get_column_letter(col)}{row}', totals.get(col), bold=col in totals, fmt=STUDDING_FORMATS.get(col))
        return row + 2

    def _section_labour_and_cost(self, ws, card, figures, row):
        merge(ws, f'A{row}:B{row}', 'Labour Details', bold=True)

        finish_rows = [(line.finish_type.name, line.rate) for line in card.finish_lines.all()]
        while len(finish_rows) < MIN_FINISH_ROWS:
            finish_rows.append(('', None))
        stone_lines = self._lines(card)
        stone_value = total(line.labour_amount for _, line in stone_lines) if stone_lines else None
        labour_rows = finish_rows + [('STONE', stone_value)]

        for offset, (label, value) in enumerate(labour_rows, start=1):
            merge(ws, f'A{row + offset}:B{row + offset}', label)
            put(ws, f'C{row + offset}', num(value))

        for col, title in zip(COST_COLUMNS, COST_HEADERS):
            put(ws, f'{col}{row}', title, bold=True)
        for col, key in zip(COST_COLUMNS, COST_KEYS):
            put(ws, f'{col}{row + 1}', num(figures[key]), fmt=MONEY)

        return row + len(labour_rows) + 2

    def _section_design(self, ws, card, figures, row):
        merge(ws, f'E{row}:F{row + 2}', 'Design Instruction')
        merge(ws, f'G{row}:P{row + 2}', card.design_note or '')
        return row + 4

    def _section_footer(self, ws, card, figures, row):
        merge(ws, f'A{row}:P{row}', self.footer_text, border=False)
        return row + 1

    @staticmethod
    def _lines(card):
        return (
            [('DIAMOND', line) for line in card.diamond_lines.all()]
            + [('COLORSTONE', line) for line in card.colorstone_lines.all()]
        )

    def _image(self, ws, path, anchor, size):
        try:
            picture = PILImage.open(path)
            picture.thumbnail(size)
            buffer = BytesIO()
            picture.save(buffer, format='PNG')
            buffer.seek(0)
            self._buffers.append(buffer)
            ws.add_image(XLImage(buffer), anchor)
        except Exception:
            pass

    def _apply_outer_border(self, ws, last_row, last_col=LAST_COL):
        for col in range(1, last_col + 1):
            top_cell = ws.cell(row=1, column=col)
            top_cell.border = Border(
                left=top_cell.border.left, right=top_cell.border.right,
                top=THIN, bottom=top_cell.border.bottom,
            )
            bottom_cell = ws.cell(row=last_row, column=col)
            bottom_cell.border = Border(
                left=bottom_cell.border.left, right=bottom_cell.border.right,
                top=bottom_cell.border.top, bottom=THIN,
            )
        for row in range(1, last_row + 1):
            left_cell = ws.cell(row=row, column=1)
            left_cell.border = Border(
                left=THIN, right=left_cell.border.right,
                top=left_cell.border.top, bottom=left_cell.border.bottom,
            )
            right_cell = ws.cell(row=row, column=last_col)
            right_cell.border = Border(
                left=right_cell.border.left, right=THIN,
                top=right_cell.border.top, bottom=right_cell.border.bottom,
            )

    def _autofit_columns(self, ws):
        merged_starts = {
            (rng.min_row, rng.min_col)
            for rng in ws.merged_cells.ranges
            if rng.max_col > rng.min_col
        }
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                if (cell.row, cell.column) in merged_starts:
                    continue
                length = len(str(cell.value))
                col = cell.column_letter
                widths[col] = max(widths.get(col, 0), length)
        for col in range(1, LAST_COL + 1):
            letter = get_column_letter(col)
            length = widths.get(letter, 0)
            ws.column_dimensions[letter].width = min(max(length, MIN_COL_WIDTH), MAX_COL_WIDTH)

    def _unique_title(self, card, used):
        base = re.sub(r'[\[\]:*?/\\]', '-', str(card.cost_card_no)).strip()[:31] or 'Sheet'
        title, count = base, 1
        while title.lower() in used:
            count += 1
            suffix = f'-{count}'
            title = base[:31 - len(suffix)] + suffix
        used.add(title.lower())
        return title