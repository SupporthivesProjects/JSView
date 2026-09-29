"""Shared openpyxl helpers for requisition exports."""

from datetime import date, datetime

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
MIN_WIDTH = 6
MAX_WIDTH = 28


def fmt_date(value):
    """Format dates as '10 Sep 2024'; pass anything else through."""
    if isinstance(value, (date, datetime)):
        return value.strftime('%d %b %Y')
    return value or ''


def num(value):
    """Decimal -> float so openpyxl writes a real number cell."""
    return None if value is None else float(value)


def put(ws, row, col, value, bold=False, border=True, fmt=None, size=11, align=CENTER):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, size=size)
    cell.alignment = align
    if border:
        cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def merge(ws, row, col, span, value, border=True, **kwargs):
    """Write a value across `span` columns and return the next free column."""
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    put(ws, row, col, value, border=border, **kwargs)
    if border:
        for c in range(col, col + span):
            ws.cell(row=row, column=c).border = BORDER
    return col + span


def fit_columns(ws, ncols):
    merged_starts = {
        (r.min_row, r.min_col) for r in ws.merged_cells.ranges if r.max_col > r.min_col
    }
    widths = {}
    for row in ws.iter_rows(max_col=ncols):
        for cell in row:
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            if (cell.row, cell.column) in merged_starts:
                continue
            longest = max(len(part) for part in str(cell.value).split('\n'))
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col in range(1, ncols + 1):
        width = min(max(widths.get(col, 0) + 2, MIN_WIDTH), MAX_WIDTH)
        ws.column_dimensions[get_column_letter(col)].width = width


def landscape_fit(ws):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)