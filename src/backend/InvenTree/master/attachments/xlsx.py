from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side

from .base import BaseAttachment, register

THIN = Side(style='thin')
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(bold=True)


class XlsxAttachment(BaseAttachment):
    mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    extension = 'xlsx'
    sheet_title = 'Sheet1'
    headers = ()

    def rows(self, ctx):
        raise NotImplementedError

    def render(self, ctx):
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet_title

        ws.append(list(self.headers))

        for cell in ws[1]:
            cell.font = HEADER_FONT

        for row in self.rows(ctx):
            ws.append(row)

        # Border only the data cells (everything below the header row).
        for data_row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(self.headers)):
            for cell in data_row:
                cell.border = CELL_BORDER

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


@register
class StoneDiamondXlsx(XlsxAttachment):
    key = 'stone_diamond_xlsx'
    sheet_title = 'Diamonds'
    headers = ('Shape', 'Size', 'Color', 'Clarity', 'Pcs', 'Weight')

    def rows(self, ctx):
        return [[d['shape'], d['size'], d['color'], d['clarity'], d['pcs'], d['weight']]
                for d in ctx['diamonds']]


@register
class StoneColorXlsx(XlsxAttachment):
    key = 'stone_color_xlsx'
    sheet_title = 'Color Stones'
    headers = ('Stone', 'Shape', 'Size', 'Colour', 'Quality', 'Pcs', 'Weight')

    def rows(self, ctx):
        return [[s['name'], s['shape'], s['size'], s['colour'], s['quality'], s['pcs'], s['weight']]
                for s in ctx['stones']]